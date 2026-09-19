#!/usr/bin/env python3
"""把岗位筛选结果（JSON）写成 xlsx，给用户在 Excel 里看和筛。
用法：python3 screen_to_xlsx.py <筛选.json> <输出.xlsx>
JSON 是一个列表，每项字段：tier(档) title cat loc nature date edu major lang pri duty why gap url [unread]；
unread 为真的行进"未读"表。纯数据表，不含公式；URL 列是可点击的超链接。"""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    sys.stdout.reconfigure(encoding='utf-8')  # Windows 终端默认不是 UTF-8，中文会乱码
except AttributeError:
    pass

COLS = [('档', 'tier', 12), ('岗位', 'title', 30), ('类别', 'cat', 12), ('地点', 'loc', 8), ('性质', 'nature', 6), ('发布日期', 'date', 11),
        ('学历', 'edu', 12), ('专业要求', 'major', 34), ('语言', 'lang', 10), ('优先项', 'pri', 30), ('工作内容一句话', 'duty', 40),
        ('理由', 'why', 44), ('缺口/剔除原因', 'gap', 34), ('链接', 'url', 50)]
TIER_ORDER = {'建议投': 0, '可投但有缺口': 1, '不建议': 2}
FILL = {'建议投': 'E2F0D9', '可投但有缺口': 'FFF2CC', '不建议': 'F2F2F2'}


def write_sheet(ws, rows, note):
    ws.append([note]); ws['A1'].font = Font(name='Arial', size=9, italic=True, color='666666')
    ws.append([c[0] for c in COLS])
    for cell in ws[2]:
        cell.font = Font(name='Arial', bold=True); cell.fill = PatternFill('solid', fgColor='D9E1F2'); cell.alignment = Alignment(vertical='top', wrap_text=True)
    for r in rows:
        ws.append([r.get(k, '') for _, k, _ in COLS])
        row = ws.max_row
        for cell in ws[row]:
            cell.font = Font(name='Arial', size=10); cell.alignment = Alignment(vertical='top', wrap_text=True)
            if r.get('tier') in FILL: cell.fill = PatternFill('solid', fgColor=FILL[r['tier']])
        url_cell = ws.cell(row=row, column=len(COLS))
        if r.get('url'): url_cell.hyperlink = r['url']; url_cell.font = Font(name='Arial', size=10, color='0563C1', underline='single')
    for i, (_, _, w) in enumerate(COLS, start=1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'C3'; ws.auto_filter.ref = f'A2:{get_column_letter(len(COLS))}{ws.max_row}'


def main(src, dst):
    rows = json.load(open(src, encoding='utf-8'))
    read = sorted([r for r in rows if not r.get('unread')], key=lambda r: (TIER_ORDER.get(r.get('tier'), 9), r.get('cat', ''), r.get('title', '')))
    unread = [r for r in rows if r.get('unread')]
    wb = Workbook(); ws = wb.active; ws.title = '筛选'
    write_sheet(ws, read, '档位与理由由模型逐份阅读 JD 得出；地点/日期/性质来自页面；用第 2 行的筛选箭头按档、类别、地点筛。')
    if unread:
        write_sheet(wb.create_sheet('未读'), unread, '按用户偏好未读详情的岗位（如实习），只有列表页信息。')
    wb.save(dst); print('saved', dst, len(read), 'read', len(unread), 'unread')


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
