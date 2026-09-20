#!/usr/bin/env python3
"""把岗位筛选结果（JSON）渲染成 md 表、xlsx 和对话里要贴的表，三处同源，列由数据决定。
用法：python3 screen_render.py <岗位筛选.json> [--md 输出.md] [--xlsx 输出.xlsx] [--title 标题] [--print 档1,档2]
  --md     写 md：读过的岗位表 + 排除清单表；文件末尾留"## 问答与定稿"供追加，表格本身不手写
  --xlsx   写 xlsx："筛选"表按档排序着色，"未读"表只有列表页信息；链接可点、有筛选箭头
  --print  把指定档位的行按 md 表打到标准输出，直接贴进对话；默认只出精简列（档、岗位、单位/部门、地点、学历、专业匹配、理由、缺口），
           --cols 列名1,列名2 可改成任意列（表头名），完整列只在 md / xlsx
JSON 是列表，每项至少有 tier、title；基础列见 BASE，其余键按第一次出现的顺序追加成列，没有的格留空、整列空也保留。
unread 为真的行进排除清单 / 未读表；以 _ 开头的键和 jobId、unread 不出列。"""
import argparse, json, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    sys.stdout.reconfigure(encoding='utf-8')  # Windows 终端默认不是 UTF-8，中文会乱码
except AttributeError:
    pass

# (表头, 键, xlsx 列宽)。基础列固定顺序；页面上读到的其他信息（招聘人数、届别、截止日期…）由 agent 写进 json 后自动追加
BASE = [('档', 'tier', 12), ('岗位', 'title', 30), ('类别', 'cat', 12), ('单位/部门', 'org', 24), ('地点', 'loc', 8), ('性质', 'nature', 6),
        ('发布日期', 'date', 11), ('学历', 'edu', 12), ('专业要求', 'major', 34), ('专业匹配', 'major_match', 12), ('语言', 'lang', 10),
        ('优先项', 'pri', 30), ('工作内容一句话', 'duty', 40), ('理由', 'why', 44), ('缺口/剔除原因', 'gap', 34), ('链接', 'url', 50)]
HIDDEN = {'unread', 'jobId'}
TIER_ORDER = {'建议投': 0, '可投但有缺口': 1, '不建议': 2}
FILL = {'建议投': 'E2F0D9', '可投但有缺口': 'FFF2CC', '不建议': 'F2F2F2'}
EXTRA_WIDTH = 14
PRINT_COLS = ['档', '岗位', '单位/部门', '地点', '学历', '专业匹配', '理由', '缺口/剔除原因']   # 对话里贴的表：宽表贴进去没人看得清


def columns(rows):
    """基础列 + 数据里多出来的键（按首次出现顺序）。"""
    cols = list(BASE)
    seen = {k for _, k, _ in BASE} | HIDDEN
    for r in rows:
        for k in r:
            if k not in seen and not k.startswith('_'):
                cols.append((k, k, EXTRA_WIDTH))
                seen.add(k)
    return cols


def split(rows):
    read = sorted([r for r in rows if not r.get('unread')], key=lambda r: (TIER_ORDER.get(r.get('tier'), 9), r.get('cat', ''), r.get('title', '')))
    unread = [r for r in rows if r.get('unread')]
    return read, unread


def cell(v):
    return '' if v is None else str(v).replace('\n', ' ').replace('|', '｜')


def md_table(rows, cols):
    lines = ['| ' + ' | '.join(h for h, _, _ in cols) + ' |', '|' + '---|' * len(cols)]
    for r in rows:
        lines.append('| ' + ' | '.join(cell(r.get(k)) for _, k, _ in cols) + ' |')
    return '\n'.join(lines)


def render_md(rows, cols, title):
    read, unread = split(rows)
    parts = [f'# {title}', '', '## 读过的岗位', md_table(read, cols), '']
    excl_cols = [('岗位', 'title', 0), ('类别', 'cat', 0), ('地点', 'loc', 0), ('剔除原因', 'gap', 0), ('链接', 'url', 0)]
    parts += ['## 排除清单（未读详情的岗位，每条带原因）', md_table(unread, excl_cols), '', '## 问答与定稿', '']
    return '\n'.join(parts)


def write_sheet(ws, rows, cols, note):
    ws.append([note]); ws['A1'].font = Font(name='Arial', size=9, italic=True, color='666666')
    ws.append([h for h, _, _ in cols])
    for c in ws[2]:
        c.font = Font(name='Arial', bold=True); c.fill = PatternFill('solid', fgColor='D9E1F2'); c.alignment = Alignment(vertical='top', wrap_text=True)
    url_col = next(i for i, (_, k, _) in enumerate(cols, start=1) if k == 'url')
    for r in rows:
        ws.append([cell(r.get(k)) for _, k, _ in cols])
        row = ws.max_row
        for c in ws[row]:
            c.font = Font(name='Arial', size=10); c.alignment = Alignment(vertical='top', wrap_text=True)
            if r.get('tier') in FILL: c.fill = PatternFill('solid', fgColor=FILL[r['tier']])
        if r.get('url'):
            u = ws.cell(row=row, column=url_col); u.hyperlink = r['url']; u.font = Font(name='Arial', size=10, color='0563C1', underline='single')
    for i, (_, _, w) in enumerate(cols, start=1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'C3'; ws.auto_filter.ref = f'A2:{get_column_letter(len(cols))}{ws.max_row}'


def render_xlsx(rows, cols, dst):
    read, unread = split(rows)
    wb = Workbook(); ws = wb.active; ws.title = '筛选'
    write_sheet(ws, read, cols, '档位、理由和硬要求列由模型逐份阅读 JD 得出；页面上没有的信息留空；用第 2 行的筛选箭头按档、类别、地点筛。')
    if unread:
        write_sheet(wb.create_sheet('未读'), unread, cols, '按用户偏好或圈定未读详情的岗位，只有列表页信息。')
    wb.save(dst)
    return len(read), len(unread)


def main(argv=None):
    ap = argparse.ArgumentParser(usage=__doc__)
    ap.add_argument('src'); ap.add_argument('--md'); ap.add_argument('--xlsx'); ap.add_argument('--title', default='岗位筛选')
    ap.add_argument('--print', dest='tiers', help='按档打印 md 表到标准输出，如 建议投,可投但有缺口')
    ap.add_argument('--cols', help='--print 时用的列，按表头名逗号分隔；不给就用精简列')
    if not (argv if argv is not None else sys.argv[1:]):
        print(__doc__); return 2
    a = ap.parse_args(argv)
    rows = json.load(open(a.src, encoding='utf-8'))
    cols = columns(rows)
    if a.md:
        with open(a.md, 'w', encoding='utf-8') as f:
            f.write(render_md(rows, cols, a.title))
        print('md', a.md)
    if a.xlsx:
        n_read, n_unread = render_xlsx(rows, cols, a.xlsx)
        print('xlsx', a.xlsx, n_read, 'read', n_unread, 'unread')
    if a.tiers:
        want = set(a.tiers.split(','))
        names = a.cols.split(',') if a.cols else PRINT_COLS
        by_name = {h: c for c in cols for h in (c[0], c[1])}
        pcols = [by_name[n] for n in names if n in by_name]
        read, _ = split(rows)
        print(md_table([r for r in read if r.get('tier') in want], pcols))
    return 0


if __name__ == '__main__':
    sys.exit(main())
