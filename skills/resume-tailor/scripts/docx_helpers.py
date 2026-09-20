# -*- coding: utf-8 -*-
"""在现有 docx 上原地改字的工具。
保留段落属性、标题横线（挂在段上的 drawing run）、项目符号定义；新 run 统一 Times New Roman / 宋体 / hint=eastAsia / 10.5 号。
用法：在投递目录的 render.py 里
    import sys; sys.path.insert(0, '<本目录>'); from docx_helpers import *
    doc = Document('复制件.docx'); P = list(doc.paragraphs)
    hdr = entry_header(P[抬头模板序号], '机构', '职位', '01/2026-04/2026')
    b = bullet(P[要点模板序号], '要点文字')
    set_text(P[标题序号], '新标题', bold=True)        # 自动保住横线
    rebuild_body(doc, [P[0], P[1], hdr, b, ...]); doc.save('out.docx')
    export_pdf_via_word('out.docx', 'out.pdf'); print(pdf_pages('out.pdf'))
"""
import copy
import os
import re
import shutil
import subprocess
import sys

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

try:
    sys.stdout.reconfigure(encoding='utf-8')  # Windows 终端默认不是 UTF-8，中文会乱码
except AttributeError:
    pass

WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"


def clear_runs(p):
    """删掉段落里所有 run 和超链接（不动 pPr）。改标题请用 set_text，它会保住横线。"""
    for r in list(p._p.findall(qn("w:r"))):
        p._p.remove(r)
    for h in list(p._p.findall(qn("w:hyperlink"))):
        p._p.remove(h)


def make_rpr(bold, ascii_font="Times New Roman", east="宋体", size_half_pt=21):
    rpr = OxmlElement("w:rPr")
    f = OxmlElement("w:rFonts")
    f.set(qn("w:ascii"), ascii_font)
    f.set(qn("w:eastAsia"), east)
    f.set(qn("w:hAnsi"), ascii_font)
    f.set(qn("w:cs"), ascii_font)
    f.set(qn("w:hint"), "eastAsia")
    rpr.append(f)
    if bold:
        rpr.append(OxmlElement("w:b"))
        rpr.append(OxmlElement("w:bCs"))
    c = OxmlElement("w:color"); c.set(qn("w:val"), "000000"); rpr.append(c)
    s = OxmlElement("w:sz"); s.set(qn("w:val"), str(size_half_pt)); rpr.append(s)
    s2 = OxmlElement("w:szCs"); s2.set(qn("w:val"), str(size_half_pt)); rpr.append(s2)
    return rpr


def add_run(p, text, bold=False, tab=False):
    r = OxmlElement("w:r")
    r.append(make_rpr(bold))
    if tab:
        r.append(OxmlElement("w:tab"))
    else:
        t = OxmlElement("w:t")
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        t.text = text
        r.append(t)
    p._p.append(r)
    return r


def set_tabs(p, mid=6000, right=10585):
    """抬头段的制表位：中间列左对齐 mid、日期右对齐 right（twips，从左页边距量起）。默认值适用 A4、左边距 1.27cm、右边距 1.09cm。"""
    ppr = p._p.get_or_add_pPr()
    for old in ppr.findall(qn("w:tabs")):
        ppr.remove(old)
    tabs = OxmlElement("w:tabs")
    for val, pos in (("left", mid), ("right", right)):
        t = OxmlElement("w:tab"); t.set(qn("w:val"), val); t.set(qn("w:pos"), str(pos)); tabs.append(t)
    # pPr 里子元素有固定顺序，w:tabs 放错位置 Word 会整个忽略：插在第一个"排在 tabs 之后"的元素前面
    after_tabs = ('w:suppressAutoHyphens', 'w:kinsoku', 'w:wordWrap', 'w:overflowPunct', 'w:topLinePunct', 'w:autoSpaceDE', 'w:autoSpaceDN',
                  'w:bidi', 'w:adjustRightInd', 'w:snapToGrid', 'w:spacing', 'w:ind', 'w:contextualSpacing', 'w:mirrorIndents',
                  'w:suppressOverlap', 'w:jc', 'w:textDirection', 'w:textAlignment', 'w:textboxTightWrap', 'w:outlineLvl', 'w:divId',
                  'w:cnfStyle', 'w:rPr', 'w:sectPr', 'w:pPrChange')
    anchor = next((c for c in ppr if c.tag in {qn(t) for t in after_tabs}), None)
    if anchor is not None:
        anchor.addprevious(tabs)
    else:
        ppr.append(tabs)
    # 首行缩进去掉，改由制表位控制
    ind = ppr.find(qn("w:ind"))
    if ind is not None:
        for a in ("w:firstLineChars", "w:firstLine"):
            if ind.get(qn(a)) is not None:
                del ind.attrib[qn(a)]


def entry_header(template, org, title, date, mid=6000, right=10585):
    """条目抬头：机构 [tab] 职位 [tab] 日期。template 为原文档中的一个抬头段落（会 deepcopy，不动原段）。"""
    p = copy.deepcopy(template._p)
    para = Paragraph(p, template._parent)
    clear_runs(para)
    set_tabs(para, mid, right)
    add_run(para, org, bold=True)
    add_run(para, "", bold=True, tab=True)
    add_run(para, title, bold=True)
    add_run(para, "", bold=True, tab=True)
    add_run(para, date, bold=True)
    return para


def bullet(template, text):
    """要点：deepcopy 模板段（带项目符号/编号定义），只换文字。"""
    p = copy.deepcopy(template._p)
    para = Paragraph(p, template._parent)
    clear_runs(para)
    add_run(para, text, bold=False)
    return para


def take_drawing_runs(para):
    """取走挂在段落上的图形 run（标题横线）。"""
    runs = [r for r in para._p.findall(qn("w:r")) if r.find(".//" + qn("w:drawing")) is not None
            or any(e.tag.endswith("}AlternateContent") for e in r.iter())]
    for r in runs:
        para._p.remove(r)
    return runs


def give_drawing_runs(para, runs, v_offset=None):
    """把图形 run 放回段落：插在 pPr 之后（放段尾会多出一行空白）。v_offset 不为 None 时改竖直偏移。"""
    for r in runs:
        for pv in r.iter("{%s}positionV" % WP):
            off = pv.find("{%s}posOffset" % WP)
            if off is not None and v_offset is not None:
                off.text = str(v_offset)
        ppr = para._p.pPr
        if ppr is not None:
            ppr.addnext(r)
        else:
            para._p.insert(0, r)


def set_text(para, text, bold=None):
    """保留段落属性和挂在段上的图形（标题横线），只换文字。"""
    keep = take_drawing_runs(para)
    clear_runs(para)
    add_run(para, text, bold=bool(bold))
    give_drawing_runs(para, keep, v_offset=None)
    return para


def rebuild_body(doc, paragraphs):
    """把正文按给定段落顺序重排（保留 sectPr）。paragraphs 是 Paragraph 对象列表，可含 deepcopy 出来的新段。
    页面设置可能挂在 body 末尾，也可能挂在最后一段的 pPr 里（那一段被删掉时页边距会回退成默认），两种都保住。"""
    body = doc.element.body
    sect = body.find(qn('w:sectPr'))
    if sect is None:
        for para in reversed(body.findall(qn('w:p'))):
            ppr = para.find(qn('w:pPr'))
            inner = ppr.find(qn('w:sectPr')) if ppr is not None else None
            if inner is not None:
                ppr.remove(inner)
                body.append(inner)
                sect = inner
                break
    for child in list(body):
        if child is not sect:
            body.remove(child)
    if sect is None:
        for para in paragraphs:
            body.append(para._p)
        return
    for para in paragraphs:
        sect.addprevious(para._p)


def export_pdf_via_word(docx_path, pdf_path):
    """用本机的 Microsoft Word 把 docx 导成 PDF，成功返回 True。
    macOS 走 AppleScript；Windows 走 docx2pdf（pip install docx2pdf，底层是 Word COM）。没有 Word 或没装依赖返回 False。"""
    docx_path, pdf_path = os.path.abspath(docx_path), os.path.abspath(pdf_path)
    if sys.platform == 'darwin':
        # Word 的 open 不返回文档对象，要用 active document
        script = f'''tell application "Microsoft Word"
  open (POSIX file "{docx_path}")
  set d to active document
  save as d file name "{pdf_path}" file format format PDF
  close d saving no
end tell'''
        r = subprocess.run(['osascript', '-e', script], capture_output=True, text=True)
        return r.returncode == 0 and os.path.exists(pdf_path)
    if sys.platform.startswith('win'):
        try:
            from docx2pdf import convert
            convert(docx_path, pdf_path)
            return os.path.exists(pdf_path)
        except Exception:
            return False
    return False


def pdf_pages(pdf_path):
    """读 PDF 页数：优先 poppler 的 pdfinfo，没有就用 pypdf；两样都没有返回 -1。"""
    if shutil.which('pdfinfo'):
        out = subprocess.run(['pdfinfo', pdf_path], capture_output=True, text=True).stdout
        m = re.search(r'Pages:\s+(\d+)', out)
        return int(m.group(1)) if m else -1
    try:
        from pypdf import PdfReader
        return len(PdfReader(pdf_path).pages)
    except Exception:
        return -1
