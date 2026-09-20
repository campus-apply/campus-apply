#!/usr/bin/env python3
"""兜底模板：用户没有现成简历时，把 resume.md 渲染成一页式中文 docx（A4，宋体 + Times New Roman 10.5 号，板块标题带底边横线，抬头用制表位对齐）。
用法：python3 render_basic.py resume.md out.docx
resume.md 格式：# 姓名 / 联系行 / ## 板块 / ### 机构 | 角色 | 起止 / - 要点"""
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

try:
    sys.stdout.reconfigure(encoding='utf-8')  # Windows 终端默认不是 UTF-8，中文会乱码
except AttributeError:
    pass


def parse(md):
    doc = {'name': '', 'contact': '', 'sections': []}
    sec = ent = None
    for raw in md.splitlines():
        line = raw.rstrip()
        if line.startswith('# ') and not doc['name']:
            doc['name'] = line[2:].strip()
        elif line.startswith('## '):
            sec = {'title': line[3:].strip(), 'entries': []}; doc['sections'].append(sec); ent = None
        elif line.startswith('### ') and sec is not None:
            parts = [x.strip() for x in line[4:].split('|')]
            ent = {'org': parts[0], 'role': parts[1] if len(parts) > 1 else '', 'date': parts[2] if len(parts) > 2 else '', 'bullets': []}
            sec['entries'].append(ent)
        elif line.startswith('- ') and sec is not None:
            if ent is None:
                ent = {'org': '', 'role': '', 'date': '', 'bullets': []}; sec['entries'].append(ent)
            ent['bullets'].append(line[2:].strip())
        elif line.strip() and doc['name'] and not doc['contact'] and sec is None:
            doc['contact'] = line.strip()
    return doc


def _font(run, bold=False, size=10.5):
    run.font.size = Pt(size); run.bold = bold
    rpr = run._r.get_or_add_rPr(); f = rpr.find(qn('w:rFonts'))
    if f is None:
        f = OxmlElement('w:rFonts'); rpr.append(f)
    for k in ('w:ascii', 'w:hAnsi', 'w:cs'):
        f.set(qn(k), 'Times New Roman')
    f.set(qn('w:eastAsia'), '宋体'); f.set(qn('w:hint'), 'eastAsia')


def _bottom_border(p):
    ppr = p._p.get_or_add_pPr(); b = OxmlElement('w:pBdr'); bt = OxmlElement('w:bottom')
    bt.set(qn('w:val'), 'single'); bt.set(qn('w:sz'), '6'); bt.set(qn('w:space'), '1'); bt.set(qn('w:color'), '1F4E79')
    b.append(bt); ppr.append(b)


def _tabs(p, mid=6000, right=10585):
    ppr = p._p.get_or_add_pPr(); tabs = OxmlElement('w:tabs')
    for val, pos in (('left', mid), ('right', right)):
        t = OxmlElement('w:tab'); t.set(qn('w:val'), val); t.set(qn('w:pos'), str(pos)); tabs.append(t)
    ppr.append(tabs)


def render(md, out_path):
    d = parse(md); doc = Document(); s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.top_margin = s.bottom_margin = Cm(1.27); s.left_margin = Cm(1.27); s.right_margin = Cm(1.09)
    st = doc.styles['Normal']; st.paragraph_format.space_after = Pt(0); st.paragraph_format.line_spacing = 1.0
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; _font(p.add_run(d['name']), bold=True, size=16)
    if d['contact']:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; _font(p.add_run(d['contact']))
    for sec in d['sections']:
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(6)
        _font(p.add_run(sec['title']), bold=True, size=11); _bottom_border(p)
        for e in sec['entries']:
            if e['org'] or e['role'] or e['date']:
                p = doc.add_paragraph(); _tabs(p)
                _font(p.add_run(e['org'] + '\t' + e['role'] + '\t' + e['date']), bold=True)
            for b in e['bullets']:
                p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.5); p.paragraph_format.first_line_indent = Cm(-0.35)
                _font(p.add_run('• ' + b))
    doc.save(out_path)


if __name__ == '__main__':
    import argparse
    ap = argparse.ArgumentParser(usage=__doc__)
    ap.add_argument('resume_md'); ap.add_argument('out_docx')
    if any(x in ('-h', '--help') for x in sys.argv[1:]):
        print(__doc__); sys.exit(0)
    if len(sys.argv) < 3:
        print(__doc__); sys.exit(2)
    a = ap.parse_args()
    render(open(a.resume_md, encoding='utf-8').read(), a.out_docx)
