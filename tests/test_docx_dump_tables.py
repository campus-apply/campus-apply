# -*- coding: utf-8 -*-
"""表格排版的简历，内容全在表格里，只扫段落会一个字都读不到。

学校和学院发的简历模板基本都是表格排版（第一列时间、第二列内容，整份简历装在一个
三十几行的表里）。只遍历 doc.paragraphs 的话，这类简历只能列出几个空段落——
改之前看不到现在写了什么，原地改字也无从下手。

这里只验"表格里的文字能被读出来"。版面占几行不在这里判——那件事靠看导出的 PDF 图，
不靠解析 docx 结构去推断（排版方式的空间太大，任何判据都只是在已知样本上调出来的）。
"""
from pathlib import Path
import subprocess
import sys

HERE = Path(__file__).resolve().parent
DUMP = HERE.parent / 'skills/resume-tailor/scripts/docx_dump.py'


def run(*args):
    return subprocess.run([sys.executable, str(DUMP), *args],
                          capture_output=True, text=True, timeout=60)


def make_table_resume(path):
    """造一份表格排版的简历：板块标题独占一行，其余行第一列时间、第二列内容。"""
    from docx import Document
    doc = Document()
    doc.add_paragraph('某某')                                # 抬头在段落里
    table = doc.add_table(rows=0, cols=2)
    for left, right in [
        ('教育背景', ''),
        ('2025.08-2027.06', '某某大学 金融学硕士'),
        ('2020.08-2025.06', '某某大学 经济学学士'),
        ('实习经历', ''),
        ('2026.05-2026.09', '某公司 产品实习生：负责一条产品线的需求拆解'),
    ]:
        row = table.add_row()
        row.cells[0].text = left
        row.cells[1].text = right
    doc.save(str(path))
    return path


def test_dump_reads_text_inside_tables(tmp_path):
    out = run(str(make_table_resume(tmp_path / 'r.docx')))
    assert out.returncode == 0, out.stdout + out.stderr
    for needle in ('教育背景', '金融学硕士', '实习经历', '需求拆解'):
        assert needle in out.stdout, '表格里的「%s」没被读出来' % needle


def test_dump_keeps_document_order(tmp_path):
    """段落和表格要按文档顺序交错输出，不能先列完段落再列表格。"""
    out = run(str(make_table_resume(tmp_path / 'r.docx')))
    text = out.stdout
    assert text.index('某某') < text.index('教育背景') < text.index('实习经历')


def test_dump_marks_table_rows_as_such(tmp_path):
    """要能看出哪一行来自表格——改字时得知道改的是单元格还是段落。"""
    out = run(str(make_table_resume(tmp_path / 'r.docx')))
    assert '表格' in out.stdout or 'TBL' in out.stdout, '表格行要标出来'


def test_merged_cells_are_not_printed_twice(tmp_path):
    """合并单元格在 python-docx 里会重复返回同一段文字，不能原样打两遍。"""
    from docx import Document
    path = tmp_path / 'm.docx'
    doc = Document()
    table = doc.add_table(rows=1, cols=3)
    merged = table.rows[0].cells[0].merge(table.rows[0].cells[2])
    merged.text = '教育背景'
    doc.save(str(path))
    out = run(str(path))
    assert out.stdout.count('教育背景') == 1, '合并单元格的文字被打了多遍'
