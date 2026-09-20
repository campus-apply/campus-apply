import importlib.util, os
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
HERE = os.path.dirname(os.path.abspath(__file__))
P = os.path.join(HERE, '..', 'skills', 'resume-tailor', 'scripts', 'docx_helpers.py')
spec = importlib.util.spec_from_file_location('docx_helpers', P); H = importlib.util.module_from_spec(spec); spec.loader.exec_module(H)


def make_doc():
    d = Document(); d.add_paragraph('标题'); d.add_paragraph('抬头 模板'); d.add_paragraph('要点模板'); return d


def test_entry_header_has_two_tabs_and_bold():
    d = make_doc(); p = H.entry_header(d.paragraphs[1], '机构', '职位', '01/2026-04/2026')
    assert p.text == '机构\t职位\t01/2026-04/2026'
    assert all(r.bold for r in p.runs if r.text)
    assert p._p.pPr.find(qn('w:tabs')) is not None


def test_bullet_clones_template_and_sets_font():
    d = make_doc(); p = H.bullet(d.paragraphs[2], '新要点')
    assert p.text == '新要点'
    rf = p.runs[0]._r.rPr.rFonts
    assert rf.get(qn('w:eastAsia')) == '宋体' and rf.get(qn('w:ascii')) == 'Times New Roman'


def test_set_text_keeps_drawing_run():
    d = make_doc(); p = d.paragraphs[0]
    r = OxmlElement('w:r'); r.append(OxmlElement('w:drawing')); p._p.append(r)
    H.set_text(p, '新标题', bold=True)
    assert p.text == '新标题' and p._p.find('.//' + qn('w:drawing')) is not None


def test_rebuild_body_reorders(tmp_path):
    d = make_doc(); ps = d.paragraphs
    H.rebuild_body(d, [ps[2], ps[0]])
    out = tmp_path / 'o.docx'; d.save(out)
    assert [p.text for p in Document(out).paragraphs] == ['要点模板', '标题']


def test_set_tabs_lands_before_run_properties_and_after_numbering_when_no_spacing():
    d = make_doc(); p = d.paragraphs[1]
    ppr = p._p.get_or_add_pPr()
    for old in list(ppr):
        ppr.remove(old)
    from docx.oxml import OxmlElement
    for tag in ('w:pStyle', 'w:numPr', 'w:ind', 'w:jc', 'w:rPr'):
        ppr.append(OxmlElement(tag))
    H.set_tabs(p)
    tags = [c.tag.split('}')[1] for c in ppr]
    assert tags.index('tabs') > tags.index('numPr') and tags.index('tabs') < tags.index('ind') < tags.index('rPr')


def test_rebuild_body_keeps_section_properties_stored_in_last_paragraph():
    d = make_doc()
    body = d.element.body
    sect = body.find(qn('w:sectPr'))
    last = d.paragraphs[-1]._p
    last.get_or_add_pPr().append(sect)          # 有些模板把 sectPr 挂在最后一段的 pPr 里
    assert body.find(qn('w:sectPr')) is None
    keep = [d.paragraphs[0], d.paragraphs[1]]
    H.rebuild_body(d, keep)
    assert body.find(qn('w:sectPr')) is not None
    assert [x.text for x in d.paragraphs[:2]] == [keep[0].text, keep[1].text]
