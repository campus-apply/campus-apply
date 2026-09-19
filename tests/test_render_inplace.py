import importlib.util, os
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.join(HERE, '..', 'skills', 'resume-tailor', 'scripts')


def load(name):
    spec = importlib.util.spec_from_file_location(name, os.path.join(SCRIPTS, name + '.py')); m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m); return m


def make_template(path):
    H = load('docx_helpers'); d = Document()
    d.add_paragraph('姓名'); d.add_paragraph('联系方式')
    for title in ('教育背景', '实习经历'):
        t = d.add_paragraph(title); r = OxmlElement('w:r'); r.append(OxmlElement('w:drawing')); t._p.append(r)
        h = d.add_paragraph('机构\t职位\t日期'); H.set_tabs(h)
        b = d.add_paragraph('要点'); ppr = b._p.get_or_add_pPr(); ppr.append(OxmlElement('w:numPr'))
    d.save(path)


MD = """# 张三
联系行

## 教育背景
### 某大学 | 硕士 | 2025-2027
- 学分绩

## 实习经历
### 某公司 | 实习生 | 2026
- 做了 A
- 做了 B
"""


def test_render_inplace_keeps_header_and_maps_sections(tmp_path):
    R = load('render_inplace'); tpl = tmp_path / 't.docx'; make_template(tpl)
    out = tmp_path / 'o.docx'; R.render(MD, str(tpl), str(out))
    texts = [p.text for p in Document(out).paragraphs]
    assert texts[:2] == ['姓名', '联系方式']
    assert texts[2] == '教育背景' and texts[3] == '某大学\t硕士\t2025-2027' and texts[4] == '学分绩'
    assert texts[5] == '实习经历' and texts[7:] == ['做了 A', '做了 B']
    d = Document(out); assert d.paragraphs[2]._p.find('.//' + qn('w:drawing')) is not None


def test_section_count_mismatch_raises(tmp_path):
    import pytest
    R = load('render_inplace'); tpl = tmp_path / 't.docx'; make_template(tpl)
    with pytest.raises(SystemExit):
        R.render("# x\n\n## 只有一个板块\n- a\n", str(tpl), str(tmp_path / 'o.docx'))
