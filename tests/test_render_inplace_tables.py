# -*- coding: utf-8 -*-
"""原地改字只处理段落排版的简历，遇到表格排版要说清为什么、该怎么办。

学校和学院发的模板基本是表格排版（整份简历装在一个三十几行的表里）。原地改字的整套逻辑
建立在"板块标题是挂横线的段落"之上，对表格无能为力——这没问题，表格要动单元格，是另一套
逻辑。问题是它失败得不清楚：报"模板里没有带制表位的抬头段"，用户看了不知道该做什么。

`docx_dump.py` 现在能列出表格内容，用户会以为能改。这个落差比两边都不行更坑，所以
失败信息必须点明"这是表格排版"。
"""
from pathlib import Path
import subprocess
import sys

HERE = Path(__file__).resolve().parent
RENDER = HERE.parent / 'skills/resume-tailor/scripts/render_inplace.py'

RESUME_MD = """# 某某

## 教育背景

### 某某大学 | 金融学硕士 | 2025.08-2027.06
- 主修课程：公司金融、资产定价
"""


def table_template(path):
    from docx import Document
    doc = Document()
    doc.add_paragraph('某某')
    table = doc.add_table(rows=0, cols=2)
    for left, right in [('教育背景', ''), ('2025.08-2027.06', '某某大学 金融学硕士')]:
        row = table.add_row()
        row.cells[0].text = left
        row.cells[1].text = right
    doc.save(str(path))
    return path


def run(md, template, out):
    return subprocess.run([sys.executable, str(RENDER), str(md), str(template), str(out)],
                          capture_output=True, text=True, timeout=60)


def test_a_table_template_is_refused_with_a_usable_reason(tmp_path):
    md = tmp_path / 'r.md'
    md.write_text(RESUME_MD, encoding='utf-8')
    r = run(md, table_template(tmp_path / 't.docx'), tmp_path / 'o.docx')
    message = r.stdout + r.stderr
    assert '表格' in message, '要点明这份模板是表格排版：' + message[:200]
    assert '原地改' in message or '另写' in message or 'render' in message, \
        '要告诉用户接下来能做什么：' + message[:200]


def test_refusing_exits_non_zero(tmp_path):
    """拒绝要退出非零，否则上游会以为改成功了。"""
    md = tmp_path / 'r.md'
    md.write_text(RESUME_MD, encoding='utf-8')
    r = run(md, table_template(tmp_path / 't.docx'), tmp_path / 'o.docx')
    assert r.returncode != 0, '拒绝了却退出 0，上游分不出成功和失败'
    assert not (tmp_path / 'o.docx').exists(), '拒绝时不该留下半成品'
