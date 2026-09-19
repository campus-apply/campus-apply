import json, os, subprocess, sys
from openpyxl import load_workbook
HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPT = os.path.join(HERE, '..', 'skills', 'job-screen', 'scripts', 'screen_render.py')

ROWS = [
    {'tier': '可投但有缺口', 'title': '甲岗', 'cat': '产品', 'org': '总部>产品部', 'loc': '深圳', 'edu': '硕士', 'major': '金融等相关', 'major_match': '本科匹配',
     'why': '有对应经历', 'gap': '硕士专业不符', 'url': 'https://jobs.example.com/1', '招聘人数': '若干', '届别': '2027 届'},
    {'tier': '建议投', 'title': '乙岗', 'cat': '产品', 'org': '总部>市场部', 'loc': '北京', 'edu': '本科及以上', 'major': '不限', 'major_match': '匹配',
     'why': '偏好相符', 'gap': '', 'url': 'https://jobs.example.com/2', '招聘人数': '2人'},
    {'tier': '未读', 'title': '丙岗', 'cat': '设计', 'loc': '上海', 'url': 'https://jobs.example.com/3', 'unread': True, 'gap': '设计类，偏好不投'},
]


def run(*args):
    return subprocess.run([sys.executable, SCRIPT, *args], capture_output=True, text=True)


def test_renders_md_and_xlsx_with_base_columns_then_extra_keys_in_order(tmp_path):
    src = tmp_path / 's.json'; src.write_text(json.dumps(ROWS, ensure_ascii=False), encoding='utf-8')
    r = run(str(src), '--md', str(tmp_path / 's.md'), '--xlsx', str(tmp_path / 's.xlsx'), '--title', '示例 岗位筛选')
    assert r.returncode == 0, r.stderr
    ws = load_workbook(tmp_path / 's.xlsx')['筛选']
    header = [c.value for c in ws[2]]
    assert header[:3] == ['档', '岗位', '类别'] and '单位/部门' in header and '专业匹配' in header
    assert header[-3:] == ['链接', '招聘人数', '届别']          # 基础列在前，多出的键按出现顺序在后
    assert [c.value for c in ws[3]][1] == '乙岗'                  # 建议投排在前
    assert ws.cell(row=3, column=header.index('届别') + 1).value in (None, '')   # 乙岗没有届别，留空
    md = (tmp_path / 's.md').read_text(encoding='utf-8')
    assert '| 招聘人数 | 届别 |' in md.splitlines()[[i for i, l in enumerate(md.splitlines()) if l.startswith('| 档 |')][0]]
    assert '## 排除清单' in md and '丙岗' in md.split('## 排除清单')[1]
    assert '若干' in md


def test_print_option_outputs_same_columns_for_chat(tmp_path):
    src = tmp_path / 's.json'; src.write_text(json.dumps(ROWS, ensure_ascii=False), encoding='utf-8')
    r = run(str(src), '--print', '建议投,可投但有缺口')
    lines = [l for l in r.stdout.splitlines() if l.startswith('|')]
    assert lines[0].startswith('| 档 | 岗位 | 类别 |') and '招聘人数' in lines[0] and '届别' in lines[0]
    assert sum('乙岗' in l or '甲岗' in l for l in lines) == 2 and not any('丙岗' in l for l in lines)


def test_missing_arguments_print_usage():
    r = run()
    assert r.returncode == 2 and '用法' in (r.stdout + r.stderr) and 'Traceback' not in r.stderr
